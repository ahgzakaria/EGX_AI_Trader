"""Regenerate the additive EGX sector reference map from the source workbook.

This writes ONE file — ``data/universe/egx_sector_map.csv`` — and nothing else.

What it deliberately does NOT do:

* it does not read, write, or validate ``data/universe/egx_universe.csv``;
* it does not change universe membership, retire a ticker, add a ticker, or
  touch Rubix subscription state;
* it does not create ``universe_changes.csv``;
* it does not attempt to resolve the known ticker-identity conflicts between
  the workbook and the canonical universe.

The map is *reference context for display*. A ticker missing from it renders as
UNKNOWN (see :mod:`core.sector_context`); it is never excluded from a signal and
never gates anything.

Regenerate after the workbook is refreshed::

    python scripts/build_egx_sector_map.py --workbook <path>

The emitted rows are sorted by ticker so the file diffs cleanly, and the source
workbook's SHA-256 is recorded on every row so a stale map is identifiable.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path
import re
import sys

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

DEFAULT_WORKBOOK = r"D:\EGX_00\EGX_Active_Stocks_By_Sector_Market_Cap.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "universe" / "egx_sector_map.csv"

#: The only file this script may write. Guards against a mistyped --output
#: silently overwriting the canonical universe.
ALLOWED_OUTPUT_NAME = "egx_sector_map.csv"

FIELDNAMES = (
    "canonical_ticker",
    "company_name",
    "sector_id",
    "sector_name",
    "sector_rank",
    "market_cap_egp",
    "source",
    "source_as_of",
    "workbook_sha256",
)


def sector_id_of(sector_name: str) -> str:
    """``Non-bank Financial Services`` -> ``non_bank_financial_services``.

    A stable join key that survives cosmetic renames of the display label.
    """

    slug = re.sub(r"[^a-z0-9]+", "_", str(sector_name).strip().lower())
    return slug.strip("_")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_workbook_rows(workbook_path: Path):
    import openpyxl

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        rows = list(workbook["All_Stocks"].iter_rows(values_only=True))
        as_of = ""
        for raw in workbook["Overview"].iter_rows(values_only=True):
            if raw and str(raw[0] or "").strip() == "Price & market cap as of":
                as_of = str(raw[1] or "").strip()
                break
    finally:
        workbook.close()

    records = []
    for raw in rows[1:]:
        ticker = str(raw[2] or "").strip().upper()
        if not ticker:
            continue
        sector_name = str(raw[12] or "").strip()
        records.append({
            "canonical_ticker": ticker,
            "company_name": str(raw[3] or "").strip(),
            "sector_id": sector_id_of(sector_name),
            "sector_name": sector_name,
            "sector_rank": raw[1],
            "market_cap_egp": raw[5],
        })
    return records, as_of


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)

    workbook_path = Path(args.workbook)
    if not workbook_path.is_file():
        parser.error(f"workbook not found: {workbook_path}")

    output = Path(args.output)
    if output.name != ALLOWED_OUTPUT_NAME:
        parser.error(
            f"refusing to write '{output.name}': this script may only write "
            f"'{ALLOWED_OUTPUT_NAME}'. The canonical universe is not this "
            "script's business."
        )

    records, as_of = read_workbook_rows(workbook_path)
    if not records:
        parser.error("workbook produced no rows; refusing to write an empty map")

    duplicates = sorted(
        {
            row["canonical_ticker"]
            for row in records
            if [r["canonical_ticker"] for r in records].count(row["canonical_ticker"]) > 1
        }
    )
    if duplicates:
        # One ticker cannot carry two sectors. Writing it would make enrichment
        # depend on row order.
        parser.error(f"duplicate tickers in workbook: {', '.join(duplicates)}")

    blank_sectors = sorted(r["canonical_ticker"] for r in records if not r["sector_name"])
    if blank_sectors:
        parser.error(f"blank sector for: {', '.join(blank_sectors)}")

    digest = _sha256(workbook_path)
    source = f"EGX sector workbook ({workbook_path.name})"
    for row in records:
        row["source"] = source
        row["source_as_of"] = as_of
        row["workbook_sha256"] = digest

    records.sort(key=lambda row: row["canonical_ticker"])
    output.parent.mkdir(parents=True, exist_ok=True)
    with open(output, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(records)

    sectors = sorted({row["sector_name"] for row in records})
    print(f"wrote {output} — {len(records)} tickers, {len(sectors)} sectors")
    print(f"workbook sha256 {digest}")
    print(f"price/market-cap as of: {as_of or '(not stated)'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
